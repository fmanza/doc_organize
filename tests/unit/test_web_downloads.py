"""Tests del Export ZIP (Fase D-2c, §2.7.4, sub-paso 2/3) y del Export Excel de caché."""

import io
import json
import zipfile

from datetime import datetime, timezone
from pathlib import Path

import httpx
import openpyxl
import pytest

from doc_organize.cache import CacheManager
from doc_organize.core.constants import DEFAULT_CACHE_DB_NAME
from doc_organize.core.models import CacheEntry, ConversionStatus, DocumentStatus, SecurityConfig
from doc_organize.projects import ProjectRegistry
from doc_organize.web import create_app


def _client(app, base_url: str = "http://testserver") -> httpx.AsyncClient:
    return httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url=base_url)


async def _create_project(tmp_path, registry_path, name: str = "Proyecto Export", tenant_id: str = "default"):
    base = tmp_path / "workspaces"
    base.mkdir(exist_ok=True)
    async with ProjectRegistry(registry_path) as registry:
        return await registry.create_project(name, base, tenant_id=tenant_id)


@pytest.mark.asyncio
async def test_export_zip_404_for_missing_project(tmp_path) -> None:
    app = create_app(registry_path=tmp_path / "projects.json")
    async with _client(app) as client:
        response = await client.get("/downloads/projects/No Existe/export.zip")
    assert response.status_code == 404


@pytest.mark.asyncio
async def test_export_zip_contains_conversion_and_cache_snapshot(tmp_path) -> None:
    registry_path = tmp_path / "projects.json"
    project = await _create_project(tmp_path, registry_path)
    root = tmp_path / "workspaces" / project.name

    conversion_dir = root / "Conversión" / "docs"
    conversion_dir.mkdir(parents=True)
    (conversion_dir / "report.md").write_text("# Informe", encoding="utf-8")

    cache_db = root / DEFAULT_CACHE_DB_NAME
    async with CacheManager(cache_db) as cache:
        await cache.save_entries(
            [
                CacheEntry(
                    file_path="/repo/docs/report.pdf",
                    sha256_hash="b" * 64,
                    timestamp=1_700_000_000.0,
                    status=DocumentStatus.NEW,
                    conversion_status=ConversionStatus.SUCCESS,
                )
            ]
        )

    app = create_app(registry_path=registry_path)
    async with _client(app) as client:
        response = await client.get(f"/downloads/projects/{project.name}/export.zip")

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"

    with zipfile.ZipFile(io.BytesIO(response.content)) as zip_file:
        names = set(zip_file.namelist())
        assert "cache_export.json" in names
        assert any(name.endswith("docs/report.md") for name in names)

        snapshot = json.loads(zip_file.read("cache_export.json"))
        assert snapshot["total_entries"] == 1
        assert snapshot["new_files"] == ["/repo/docs/report.pdf"]


@pytest.mark.asyncio
async def test_export_zip_includes_dedup_report_when_present(tmp_path) -> None:
    registry_path = tmp_path / "projects.json"
    project = await _create_project(tmp_path, registry_path)
    root = tmp_path / "workspaces" / project.name

    duplicados_dir = root / "duplicados"
    duplicados_dir.mkdir(parents=True)
    (duplicados_dir / "DEDUP_REPORT.json").write_text(json.dumps({"groups": []}), encoding="utf-8")

    app = create_app(registry_path=registry_path)
    async with _client(app) as client:
        response = await client.get(f"/downloads/projects/{project.name}/export.zip")

    with zipfile.ZipFile(io.BytesIO(response.content)) as zip_file:
        assert any(name.endswith("duplicados/DEDUP_REPORT.json") for name in zip_file.namelist())


@pytest.mark.asyncio
async def test_export_zip_works_when_never_scanned(tmp_path) -> None:
    # Sin Conversion/ ni cache.db previos: el ZIP igual se genera, solo con cache_export.json vacio.
    registry_path = tmp_path / "projects.json"
    project = await _create_project(tmp_path, registry_path)

    app = create_app(registry_path=registry_path)
    async with _client(app) as client:
        response = await client.get(f"/downloads/projects/{project.name}/export.zip")

    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.content)) as zip_file:
        snapshot = json.loads(zip_file.read("cache_export.json"))
        assert snapshot["total_entries"] == 0


@pytest.mark.asyncio
async def test_export_zip_tenant_isolation(tmp_path) -> None:
    registry_path = tmp_path / "projects.json"
    project = await _create_project(tmp_path, registry_path, tenant_id="acme-corp")

    # Loopback (sin sesion) -> get_tenant_id() resuelve "default", distinto del tenant real.
    app = create_app(registry_path=registry_path)
    async with _client(app) as client:
        response = await client.get(f"/downloads/projects/{project.name}/export.zip")

    assert response.status_code == 404


@pytest.mark.asyncio
async def test_export_zip_requires_session_when_non_loopback(tmp_path) -> None:
    registry_path = tmp_path / "projects.json"
    project = await _create_project(tmp_path, registry_path)
    security = SecurityConfig(bind_host="0.0.0.0")
    app = create_app(security=security, registry_path=registry_path, auth_db_path=tmp_path / "auth.db")

    async with _client(app) as client:
        response = await client.get(f"/downloads/projects/{project.name}/export.zip")

    assert response.status_code == 401


def _read_xlsx_rows(content: bytes) -> list[tuple]:
    workbook = openpyxl.load_workbook(io.BytesIO(content))
    sheet = workbook["Cache"]
    return list(sheet.iter_rows(values_only=True))


@pytest.mark.asyncio
async def test_export_cache_xlsx_404_for_missing_project(tmp_path) -> None:
    app = create_app(registry_path=tmp_path / "projects.json")
    async with _client(app) as client:
        response = await client.get("/downloads/projects/No Existe/cache-export.xlsx")
    assert response.status_code == 404


@pytest.mark.asyncio
async def test_export_cache_xlsx_works_when_never_scanned(tmp_path) -> None:
    # Sin cache.db previo: se genera igualmente un .xlsx valido, solo con cabeceras.
    registry_path = tmp_path / "projects.json"
    project = await _create_project(tmp_path, registry_path)

    app = create_app(registry_path=registry_path)
    async with _client(app) as client:
        response = await client.get(f"/downloads/projects/{project.name}/cache-export.xlsx")

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    rows = _read_xlsx_rows(response.content)
    assert rows == [("Archivo", "Carpeta", "Ruta completa", "Hash SHA-256", "Estado", "Conversión", "Última actualización (UTC)")]


@pytest.mark.asyncio
async def test_export_cache_xlsx_contains_full_hash_and_all_entries(tmp_path) -> None:
    registry_path = tmp_path / "projects.json"
    project = await _create_project(tmp_path, registry_path)
    cache_db = tmp_path / "workspaces" / project.name / DEFAULT_CACHE_DB_NAME
    full_hash = "b" * 64
    async with CacheManager(cache_db) as cache:
        await cache.save_entries(
            [
                CacheEntry(
                    file_path="/repo/docs/report.pdf",
                    sha256_hash=full_hash,
                    timestamp=1_700_000_000.0,
                    status=DocumentStatus.NEW,
                    conversion_status=ConversionStatus.SUCCESS,
                ),
                CacheEntry(
                    file_path="/repo/docs/other.pdf",
                    sha256_hash="c" * 64,
                    timestamp=1_700_000_100.0,
                    status=DocumentStatus.UNCHANGED,
                    conversion_status=None,
                ),
            ]
        )

    app = create_app(registry_path=registry_path)
    async with _client(app) as client:
        response = await client.get(f"/downloads/projects/{project.name}/cache-export.xlsx")

    assert response.status_code == 200
    rows = _read_xlsx_rows(response.content)
    assert len(rows) == 3  # cabecera + 2 entradas
    # Orden estable por file_path: other.pdf antes que report.pdf.
    header, first, second = rows
    assert first[0] == "other.pdf"
    # `Path("/repo/docs/other.pdf").parent` normaliza a separadores nativos (p. ej.
    # backslash en Windows) -- se compara contra Path en vez de una cadena literal.
    assert first[1] == str(Path("/repo/docs"))
    assert first[2] == "/repo/docs/other.pdf"
    assert first[3] == "c" * 64
    assert first[4] == "unchanged"
    assert first[5] is None  # celda vacia: openpyxl la lee como None, no como ""
    assert second[0] == "report.pdf"
    assert second[3] == full_hash  # hash completo, no truncado a 12 caracteres
    assert second[4] == "new"
    assert second[5] == "success"
    assert second[6] == datetime.fromtimestamp(1_700_000_000.0, tz=timezone.utc).isoformat()


@pytest.mark.asyncio
async def test_export_cache_xlsx_neutralizes_formula_injection_in_file_name(tmp_path) -> None:
    # Revision de seguridad: un fichero con un nombre que empieza por "=" se
    # interpretaria como formula al abrir el .xlsx en Excel/LibreOffice si se
    # escribe tal cual -- debe anteponerse un "'" (mitigacion OWASP CSV/Excel
    # Formula Injection). Sin "/" ni "\" en el nombre: Windows los trataria como
    # separador de ruta y `Path(...).name` se quedaria solo con lo que va detras.
    registry_path = tmp_path / "projects.json"
    project = await _create_project(tmp_path, registry_path)
    cache_db = tmp_path / "workspaces" / project.name / DEFAULT_CACHE_DB_NAME
    malicious_name = "=cmd|'calc'!A1.pdf"
    async with CacheManager(cache_db) as cache:
        await cache.save_entries(
            [
                CacheEntry(
                    file_path=f"/repo/{malicious_name}",
                    sha256_hash="a" * 64,
                    timestamp=1_700_000_000.0,
                    status=DocumentStatus.NEW,
                    conversion_status=None,
                )
            ]
        )

    app = create_app(registry_path=registry_path)
    async with _client(app) as client:
        response = await client.get(f"/downloads/projects/{project.name}/cache-export.xlsx")

    rows = _read_xlsx_rows(response.content)
    _, data_row = rows
    # "Archivo" (solo el nombre) empieza por "=" -> se neutraliza con un "'" inicial.
    assert data_row[0] == f"'{malicious_name}"
    # "Ruta completa" empieza por "/repo/..." (no por un caracter disparador de
    # formula) -> no hace falta neutralizarla, se conserva tal cual.
    assert data_row[2] == f"/repo/{malicious_name}"


@pytest.mark.asyncio
async def test_export_cache_xlsx_tenant_isolation(tmp_path) -> None:
    registry_path = tmp_path / "projects.json"
    project = await _create_project(tmp_path, registry_path, tenant_id="acme-corp")

    app = create_app(registry_path=registry_path)
    async with _client(app) as client:
        response = await client.get(f"/downloads/projects/{project.name}/cache-export.xlsx")

    assert response.status_code == 404


@pytest.mark.asyncio
async def test_export_cache_xlsx_requires_session_when_non_loopback(tmp_path) -> None:
    registry_path = tmp_path / "projects.json"
    project = await _create_project(tmp_path, registry_path)
    security = SecurityConfig(bind_host="0.0.0.0")
    app = create_app(security=security, registry_path=registry_path, auth_db_path=tmp_path / "auth.db")

    async with _client(app) as client:
        response = await client.get(f"/downloads/projects/{project.name}/cache-export.xlsx")

    assert response.status_code == 401


def _read_xlsx_sheet_rows(content: bytes, sheet_name: str) -> list[tuple]:
    workbook = openpyxl.load_workbook(io.BytesIO(content))
    sheet = workbook[sheet_name]
    return list(sheet.iter_rows(values_only=True))


@pytest.mark.asyncio
async def test_export_scan_errors_xlsx_404_for_missing_project(tmp_path) -> None:
    app = create_app(registry_path=tmp_path / "projects.json")
    async with _client(app) as client:
        response = await client.get("/downloads/projects/No Existe/scan-errors-export.xlsx")
    assert response.status_code == 404


@pytest.mark.asyncio
async def test_export_scan_errors_xlsx_has_current_and_history_sheets(tmp_path) -> None:
    registry_path = tmp_path / "projects.json"
    project = await _create_project(tmp_path, registry_path)
    project_root = tmp_path / "workspaces" / project.name

    cache_db = project_root / DEFAULT_CACHE_DB_NAME
    async with CacheManager(cache_db) as cache:
        await cache.save_entries(
            [
                CacheEntry(
                    file_path="/repo/docs/roto.pdf",
                    sha256_hash="a" * 64,
                    timestamp=1_700_000_000.0,
                    status=DocumentStatus.NEW,
                    conversion_status=ConversionStatus.ERROR,
                    error_message="Documento cifrado",
                    error_retry_count=1,
                )
            ]
        )

    history_dir = project_root / "errores_escaneo"
    history_dir.mkdir(parents=True)
    report = {
        "scan_run_id": "run-1",
        "timestamp": "2026-08-01T10:00:00+00:00",
        "total_scanned": 5,
        "entries": [{"file_path": "/repo/docs/viejo.pdf", "error_message": "Error previo", "error_retry_count": 3, "capped": True}],
    }
    (history_dir / "SCAN_ERRORS_HISTORY.jsonl").write_text(json.dumps(report) + "\n", encoding="utf-8")

    app = create_app(registry_path=registry_path)
    async with _client(app) as client:
        response = await client.get(f"/downloads/projects/{project.name}/scan-errors-export.xlsx")

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Estado actual", "Histórico"]

    current_rows = _read_xlsx_sheet_rows(response.content, "Estado actual")
    assert current_rows[0] == ("Archivo", "Carpeta", "Ruta completa", "Mensaje de error", "Reintentos", "Límite alcanzado")
    assert current_rows[1][0] == "roto.pdf"
    assert current_rows[1][3] == "Documento cifrado"
    assert current_rows[1][5] == "No"

    history_rows = _read_xlsx_sheet_rows(response.content, "Histórico")
    assert history_rows[0] == (
        "Ejecución (scan_run_id)",
        "Fecha (UTC)",
        "Archivo",
        "Carpeta",
        "Ruta completa",
        "Mensaje de error",
        "Reintentos",
        "Límite alcanzado",
    )
    assert history_rows[1][0] == "run-1"
    assert history_rows[1][2] == "viejo.pdf"
    assert history_rows[1][7] == "Sí"


@pytest.mark.asyncio
async def test_export_scan_errors_xlsx_empty_state_has_headers_only(tmp_path) -> None:
    registry_path = tmp_path / "projects.json"
    project = await _create_project(tmp_path, registry_path)

    app = create_app(registry_path=registry_path)
    async with _client(app) as client:
        response = await client.get(f"/downloads/projects/{project.name}/scan-errors-export.xlsx")

    assert response.status_code == 200
    assert _read_xlsx_sheet_rows(response.content, "Estado actual") == [
        ("Archivo", "Carpeta", "Ruta completa", "Mensaje de error", "Reintentos", "Límite alcanzado")
    ]
    assert _read_xlsx_sheet_rows(response.content, "Histórico") == [
        (
            "Ejecución (scan_run_id)",
            "Fecha (UTC)",
            "Archivo",
            "Carpeta",
            "Ruta completa",
            "Mensaje de error",
            "Reintentos",
            "Límite alcanzado",
        )
    ]


@pytest.mark.asyncio
async def test_export_scan_errors_xlsx_tenant_isolation(tmp_path) -> None:
    registry_path = tmp_path / "projects.json"
    project = await _create_project(tmp_path, registry_path, tenant_id="acme-corp")

    app = create_app(registry_path=registry_path)
    async with _client(app) as client:
        response = await client.get(f"/downloads/projects/{project.name}/scan-errors-export.xlsx")

    assert response.status_code == 404
